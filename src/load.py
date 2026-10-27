"""Build citibike.duckdb: run sql/schema.sql, then load every staged source into the ERD tables.

pandas DataFrames are inserted with DuckDB replacement scans: `INSERT ... SELECT ... FROM days` reads the local
DataFrame named `days`. Frames that are not otherwise used in Python are registered explicitly (con.register).
"""
import re
import warnings
from pathlib import Path

import duckdb
import openpyxl
import pandas as pd
from pandas.tseries.holiday import USFederalHolidayCalendar

ROOT = Path(__file__).resolve().parents[1]
STAGING = ROOT / "data" / "staging"
DB = ROOT / "citibike.duckdb"
START, END = "2023-07-01", "2026-06-30"
SAMPLE_MONTHS = ["202507", "202510", "202601", "202604"]

CITIES = [  # city_id, city, system, operator, NOAA station
    ("NYC", "New York City", "Citi Bike", "Lyft (NYC Bike Share LLC)", "USW00094728"),
    ("CHI", "Chicago", "Divvy", "Lyft", "USW00094846"),
]
PLANS = [  # plan_id, city, name, billing, subscription?, rider_type
    ("NYC_SINGLE", "NYC", "Single Ride", "per ride", False, "casual"),
    ("NYC_DAY", "NYC", "Day Pass", "day", False, "casual"),
    ("NYC_THREEDAY", "NYC", "Three-Day Pass", "three days", False, "casual"),
    ("NYC_ANNUAL", "NYC", "Annual Membership", "year", True, "member"),
    ("NYC_LYFTPINK", "NYC", "Lyft Pink", "year", True, "member"),
    ("NYC_LYFTPINK_MONTHLY", "NYC", "Lyft Pink (monthly plan)", "month", True, "member"),
    ("NYC_EMPLOYER", "NYC", "Discounted Annual (employer/school)", "year", True, "member"),
    ("NYC_REDUCED_FARE", "NYC", "Reduced Fare Bike Share", "month", True, "member"),
    ("CHI_SINGLE", "CHI", "Single Ride", "per ride", False, "casual"),
    ("CHI_DAY", "CHI", "Day Pass", "day", False, "casual"),
    ("CHI_ANNUAL", "CHI", "Annual Membership", "year", True, "member"),
    ("CHI_ANNUAL_NEW_OFFER", "CHI", "Annual Membership (new-member offer)", "year", True, "member"),
    ("CHI_LYFTPINK", "CHI", "Lyft Pink", "year", True, "member"),
]
PRICE_PLAN = {"Single Ride": "SINGLE", "Day Pass": "DAY", "Annual Membership": "ANNUAL",
              "Lyft Pink": "LYFTPINK", "Annual Membership (new-member offer)": "ANNUAL_NEW_OFFER"}
PASS_PLAN = {"annual": "NYC_ANNUAL", "single-trip": "NYC_SINGLE", "single-day": "NYC_DAY",
             "three-day": "NYC_THREEDAY"}
SURVEY_ACCESS_PLAN = {1: "NYC_SINGLE", 2: "NYC_ANNUAL", 3: "NYC_LYFTPINK", 4: "NYC_EMPLOYER",
                      5: "NYC_REDUCED_FARE", 6: "NYC_LYFTPINK_MONTHLY", 7: "NYC_DAY"}
SURVEY_VARS = ["age", "gender", "employment", "citi_bike_freq", "citi_bike_type", "citi_bike_pay"] + \
              [f"no_bike_share_{i}" for i in range(1, 9)]


def season(month: int) -> str:
    return {12: "Winter", 1: "Winter", 2: "Winter", 3: "Spring", 4: "Spring", 5: "Spring",
            6: "Summer", 7: "Summer", 8: "Summer"}.get(month, "Autumn")


def survey_codes() -> pd.DataFrame:
    """Value labels parsed from the official CMS 2024 data dictionary (Column Information sheet)."""
    with warnings.catch_warnings():
        # openpyxl warns that it cannot read the sheet's data-validation extension; the values are unaffected
        warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
        wb = openpyxl.load_workbook(ROOT / "data/raw/survey/cms_2024_data_dictionary.xlsx", read_only=True)
        sheet = list(wb["Column Information"].iter_rows(values_only=True))
    rows = []
    for r in sheet:
        if r[0] == "Person" and r[1] in SURVEY_VARS:
            values = next((c for c in r[4:] if isinstance(c, str) and re.search(r"^\d+-", c, re.M)), "")
            for code, label in re.findall(r"^(\d+)-\s*(.+)$", values, re.M):
                rows.append((r[1], int(code), label.strip()))
            if r[1].startswith("no_bike_share_"):
                # the reason text is in the column description: "Reason for not using Citi Bike: <reason>"
                reason = str(r[3]).split(":", 1)[1].strip()
                rows.append(("non_use_reason", int(r[1].rsplit("_", 1)[1]), reason))
    return pd.DataFrame(rows, columns=["variable", "code", "label"])


def main() -> None:
    DB.unlink(missing_ok=True)
    con = duckdb.connect(str(DB))
    con.execute((ROOT / "sql" / "schema.sql").read_text())

    con.executemany("INSERT INTO city VALUES (?, ?, ?, ?, ?)", CITIES)
    con.executemany("INSERT INTO rider_type VALUES (?, ?)", [
        ("member", "Annual or monthly member (subscription)"),
        ("casual", "Single ride or day/multi-day pass (pay as you go)")])

    days = pd.DataFrame({"date": pd.date_range(START, END, freq="D")})
    holidays = set(USFederalHolidayCalendar().holidays(START, END))
    days = days.assign(year=days.date.dt.year, month=days.date.dt.month,
                       year_month=days.date.dt.strftime("%Y-%m"), day_of_week=days.date.dt.day_name(),
                       is_weekend=days.date.dt.dayofweek >= 5, is_holiday=days.date.isin(holidays),
                       season=days.date.dt.month.map(season))
    con.execute("INSERT INTO calendar_date SELECT * FROM days")

    con.executemany("INSERT INTO plan VALUES (?, ?, ?, ?, ?, ?)", PLANS)
    prices = pd.read_csv(STAGING / "plan_price_observations.csv")
    prices["plan_id"] = prices.city + "_" + prices.plan.map(PRICE_PLAN)
    prices = prices.drop_duplicates(["plan_id", "observed_date"])
    con.execute("""INSERT INTO plan_price SELECT plan_id, observed_date::DATE, price_usd, per_minute_usd, source_url
                   FROM prices""")

    con.execute(f"""INSERT INTO weather_daily
        SELECT city, date::DATE, tmax_c, tmin_c, precip_mm, snow_mm, snow_depth_mm, wind_ms
        FROM read_csv('{STAGING / "weather_daily.csv"}')""")
    bts = f"read_csv('{STAGING / 'bts_daily.csv'}')"
    con.execute(f"INSERT INTO qa_bts_daily SELECT city, date::DATE, rider_type, trips::BIGINT, trip_minutes FROM {bts}")
    con.execute(f"INSERT INTO qa_bts_jersey_city_daily SELECT date::DATE, rider_type, trips::BIGINT "
                f"FROM read_csv('{STAGING / 'bts_daily_jersey_city.csv'}')")
    # NYC: BTS (matches Citi Bike files within 0.1%). CHI: Divvy's own files for all 36 months, counted with
    # the BTS rule (end date, 1-120 min; Divvy allows parking outside stations so no station requirement).
    con.execute(f"INSERT INTO daily_ridership SELECT city, date::DATE, rider_type, trips::BIGINT, trip_minutes, "
                f"'BTS' FROM {bts} WHERE city = 'NYC'")
    con.execute(f"""INSERT INTO daily_ridership
        SELECT 'CHI', ended_at::DATE, member_casual, count(*),
               sum(date_diff('second', started_at, ended_at)) / 60.0, 'Divvy trip files'
        FROM read_parquet('{STAGING / "trips"}/CHI_*.parquet')
        WHERE date_diff('second', started_at, ended_at) BETWEEN 60 AND 7200
          AND ended_at::DATE BETWEEN '{START}' AND '{END}'
        GROUP BY ALL""")

    ops = pd.read_csv(STAGING / "nyc_monthly_operations.csv")
    ops["report_published"] = ops.source_file.notna()
    # scope change found by QA: report trip totals include Jersey City/Hoboken from Mar 2025
    ops["trips_include_jersey_city"] = ops.report_published & (ops.month >= "2025-03")
    con.execute("""INSERT INTO monthly_operations
        SELECT 'NYC', month, report_published, total_annual_members, report_mentions_jersey_city,
               trips_include_jersey_city,
               annual_signups_or_renewals, trips_total, trips_member, trips_casual, avg_trip_minutes,
               active_stations, avg_fleet, source_file FROM ops""")
    con.execute("""INSERT INTO support_contacts
        SELECT 'NYC', month, 'phone', cs_calls, cs_avg_handle_minutes, source_file FROM ops WHERE report_published
        UNION ALL
        SELECT 'NYC', month, 'email/web form', cs_emails_webforms, NULL, source_file FROM ops WHERE report_published""")
    passes = pd.read_csv(STAGING / "nyc_pass_sales.csv")
    passes["plan_id"] = passes.pass_type.map(PASS_PLAN)
    con.execute("INSERT INTO pass_sales SELECT plan_id, month, quantity, source_file FROM passes")
    con.register("rev", pd.read_csv(STAGING / "nyc_revenue.csv"))
    con.execute("INSERT INTO revenue SELECT 'NYC', month, category, amount_mtd_usd, amount_ytd_usd, basis, "
                "source_file FROM rev")

    trips = "[" + ", ".join(f"'{STAGING / 'trips'}/{c}_{ym}.parquet'"
                            for c in ("NYC", "CHI") for ym in SAMPLE_MONTHS) + "]"
    con.execute(f"""INSERT INTO station
        SELECT city || ':' || station_id, city, station_id, mode(station_name), avg(lat), avg(lng)
        FROM (SELECT city, start_station_id AS station_id, start_station_name AS station_name,
                     start_lat AS lat, start_lng AS lng FROM read_parquet({trips})
              UNION ALL
              SELECT city, end_station_id, end_station_name, end_lat, end_lng FROM read_parquet({trips}))
        WHERE station_id IS NOT NULL GROUP BY city, station_id""")
    con.execute(f"""INSERT INTO trip_sample
        SELECT city, ride_id, member_casual, rideable_type, started_at, ended_at, started_at::DATE,
               date_diff('second', started_at, ended_at) / 60.0 AS duration_min,
               CASE WHEN start_station_id IS NOT NULL THEN city || ':' || start_station_id END,
               CASE WHEN end_station_id IS NOT NULL THEN city || ':' || end_station_id END,
               (date_diff('second', started_at, ended_at) BETWEEN 60 AND 7200)
                 AND (city = 'CHI' OR (start_station_id IS NOT NULL AND end_station_id IS NOT NULL))
        FROM read_parquet({trips})
        WHERE started_at::DATE BETWEEN '{START}' AND '{END}'""")

    con.register("codes", survey_codes())
    con.execute("INSERT INTO survey_code SELECT * FROM codes")
    p = pd.read_csv(STAGING / "survey_person_2024.csv")
    con.execute("""INSERT INTO survey_respondent
        SELECT person_id, hh_id, age, gender, employment, citi_bike_freq, citi_bike_type, person_weight_citywide
        FROM p""")
    held = p.loc[~p.citi_bike_pay.isin([995]), ["person_id", "citi_bike_pay"]].copy()
    held["plan_id"] = held.citi_bike_pay.map(SURVEY_ACCESS_PLAN)
    con.execute("INSERT INTO survey_plan_held SELECT person_id, citi_bike_pay, plan_id FROM held")
    reasons = (p.melt(id_vars="person_id", value_vars=[f"no_bike_share_{i}" for i in range(1, 9)])
                .query("value == 1"))
    reasons["reason_code"] = reasons.variable.str.extract(r"(\d+)$").astype(int)
    con.execute("INSERT INTO survey_non_use_reason SELECT person_id, reason_code FROM reasons")

    for (t,) in con.execute("SELECT table_name FROM information_schema.tables ORDER BY 1").fetchall():
        print(f"{t:22s} {con.execute(f'SELECT count(*) FROM {t}').fetchone()[0]:>12,}")
    con.close()


if __name__ == "__main__":
    main()
